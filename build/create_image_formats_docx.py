from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/hoangminhan/cv_profile/build/Bang_so_sanh_dinh_dang_anh.docx")


FORMATS = [
    (
        "PNG",
        "Ảnh điểm ảnh\nNén không mất dữ liệu\nCó nền trong suốt",
        "Giữ chữ, đường nét và màu phẳng rất rõ. Lưu lại nhiều lần ít làm giảm chất lượng.",
        "Thường nặng hơn JPG khi dùng cho ảnh chụp.",
        "Logo, icon, ảnh chụp màn hình, sơ đồ và ảnh cần nền trong suốt.",
    ),
    (
        "JPG\nJPEG",
        "Ảnh điểm ảnh\nNén mất dữ liệu\nKhông có nền trong suốt",
        "Dung lượng nhỏ, tương thích rộng và phù hợp với ảnh có nhiều màu sắc.",
        "Chữ và viền dễ bị nhòe khi nén mạnh. Chất lượng giảm nếu lưu lại nhiều lần.",
        "Ảnh chân dung, phong cảnh, sản phẩm, email và mạng xã hội.",
    ),
    (
        "WebP",
        "Ảnh điểm ảnh\nNén mất hoặc không mất dữ liệu\nCó trong suốt và ảnh động",
        "Thường nhẹ hơn JPG và PNG ở chất lượng tương đương; dùng linh hoạt trên web.",
        "Một số phần mềm hoặc quy trình cũ hỗ trợ chưa tốt.",
        "Website, cửa hàng trực tuyến và ứng dụng cần ảnh nhẹ.",
    ),
    (
        "AVIF",
        "Ảnh điểm ảnh\nNén hiệu quả cao\nCó trong suốt và HDR",
        "Dung lượng rất nhỏ nhưng vẫn giữ chi tiết và màu sắc tốt.",
        "Mã hóa có thể chậm; phần mềm cũ có thể không mở được.",
        "Website hiện đại, ảnh HDR và hệ thống ưu tiên tiết kiệm băng thông.",
    ),
    (
        "GIF",
        "Ảnh điểm ảnh\nTối đa 256 màu mỗi khung\nCó ảnh động",
        "Dễ chia sẻ và tự chạy lặp trên nhiều nền tảng.",
        "Màu hạn chế; ảnh động dài thường nặng hơn video hoặc WebP.",
        "Meme, sticker, biểu tượng động và hướng dẫn chuyển động ngắn.",
    ),
    (
        "SVG",
        "Đồ họa vector\nDựa trên XML\nCó thể trong suốt và tạo chuyển động",
        "Phóng lớn không vỡ, thường nhẹ với hình đơn giản và dễ đổi màu.",
        "Không phù hợp với ảnh chụp; cần kiểm soát nội dung SVG khi lấy từ nguồn lạ.",
        "Logo, icon, biểu đồ, sơ đồ và hình minh họa trên web.",
    ),
    (
        "HEIC\nHEIF",
        "Định dạng chứa ảnh hiệu quả cao\nThường dùng trên iPhone",
        "Chất lượng tốt với dung lượng nhỏ; có thể lưu nhiều ảnh và thông tin bổ sung.",
        "Khả năng tương thích với Windows, website và phần mềm cũ còn hạn chế.",
        "Lưu ảnh trên thiết bị Apple; chuyển sang JPG khi cần chia sẻ rộng.",
    ),
    (
        "TIFF",
        "Ảnh điểm ảnh\nCó thể nén không mất dữ liệu\nHỗ trợ độ sâu màu cao",
        "Giữ chất lượng rất cao, phù hợp với lưu trữ, scan và quy trình in chuyên nghiệp.",
        "Dung lượng lớn và không thích hợp để đăng web thông thường.",
        "In ấn, xuất bản, scan tài liệu và lưu trữ ảnh chất lượng cao.",
    ),
    (
        "RAW",
        "Dữ liệu thô từ cảm biến\nVí dụ CR3, NEF, ARW và DNG",
        "Giữ nhiều dữ liệu sáng và màu, cho phép hậu kỳ linh hoạt.",
        "File lớn, cần phần mềm xử lý và thường phải xuất sang định dạng khác để chia sẻ.",
        "Chụp ảnh chuyên nghiệp và các ảnh cần chỉnh màu, sáng hoặc vùng tối sâu.",
    ),
    (
        "BMP",
        "Ảnh bitmap\nThường ít nén hoặc không nén",
        "Cấu trúc đơn giản và được một số hệ thống cũ hỗ trợ.",
        "Dung lượng rất lớn, ít tính năng hiện đại và không tối ưu cho web.",
        "Phần mềm cũ hoặc trường hợp kỹ thuật đặc biệt yêu cầu BMP.",
    ),
    (
        "PSD",
        "File làm việc của Photoshop\nGiữ layer, chữ và hiệu ứng",
        "Giữ nguyên cấu trúc thiết kế để tiếp tục chỉnh sửa.",
        "Không phù hợp làm file giao cuối; thường cần Photoshop hoặc phần mềm tương thích.",
        "Thiết kế và chỉnh sửa ảnh nhiều lớp trước khi xuất PNG, JPG hoặc PDF.",
    ),
    (
        "AI",
        "File vector của Adobe Illustrator\nGiữ layer và đối tượng",
        "Tốt cho thiết kế vector có thể thay đổi kích thước và chỉnh sửa.",
        "Phụ thuộc phần mềm tương thích; không nên dùng trực tiếp làm ảnh chia sẻ phổ thông.",
        "Logo, bao bì, minh họa vector và tài liệu thiết kế đang chỉnh sửa.",
    ),
    (
        "XCF",
        "File làm việc của GIMP\nGiữ layer và vùng chọn",
        "Phù hợp để tiếp tục chỉnh sửa trong GIMP và các quy trình nguồn mở.",
        "Khả năng tương thích ngoài GIMP hạn chế; cần xuất sang định dạng phổ biến.",
        "Lưu bản thiết kế gốc khi chỉnh sửa ảnh bằng GIMP.",
    ),
    (
        "KRA",
        "File làm việc của Krita\nGiữ layer và cọ vẽ",
        "Giữ đầy đủ dữ liệu phục vụ vẽ và chỉnh sửa tranh kỹ thuật số.",
        "Ít phần mềm khác hỗ trợ; không phải định dạng giao thành phẩm.",
        "Minh họa, truyện tranh và hội họa kỹ thuật số trong Krita.",
    ),
    (
        "ICO",
        "Tệp biểu tượng\nCó thể chứa nhiều kích thước\nCó nền trong suốt",
        "Gói nhiều phiên bản icon trong một file để hệ thống chọn kích thước phù hợp.",
        "Chỉ phù hợp với biểu tượng, không dùng cho ảnh thông thường.",
        "Icon ứng dụng Windows, shortcut và một số favicon.",
    ),
    (
        "PDF",
        "Tài liệu cố định bố cục\nCó thể chứa ảnh điểm ảnh và vector",
        "Giữ bố cục, chữ và hình ảnh ổn định khi gửi, xem hoặc in.",
        "Không phải file ảnh để chỉnh sửa trực tiếp; ảnh bên trong vẫn có thể bị nén.",
        "CV, tài liệu hoàn chỉnh, bản in và file gửi cho người khác xem.",
    ),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=105, bottom=95, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(run, size=9, color="666666")


def configure_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title._element.find(qn("w:pPr"))
    if title_p_pr is not None:
        title_border = title_p_pr.find(qn("w:pBdr"))
        if title_border is not None:
            title_p_pr.remove(title_border)

    heading = styles["Heading 1"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading.font.size = Pt(15)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True


def create_document():
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    configure_styles(document)

    document.core_properties.title = "Bảng so sánh các định dạng ảnh"
    document.core_properties.subject = "Bảng tra cứu PNG JPG WebP AVIF GIF SVG HEIC TIFF RAW và các định dạng thiết kế"
    document.core_properties.author = ""
    document.core_properties.keywords = "định dạng ảnh, PNG, JPG, WebP, AVIF, GIF, SVG, HEIC, TIFF, RAW"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Bảng so sánh các định dạng ảnh")

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(9)
    intro.add_run(
        "Tài liệu này giúp chọn định dạng phù hợp để lưu, chỉnh sửa, chia sẻ, đăng web hoặc in ảnh. "
        "Ảnh chụp thường phù hợp với JPG, WebP hoặc AVIF; logo và hình có nền trong suốt thường phù hợp với PNG hoặc SVG; "
        "các file RAW, PSD, AI, XCF và KRA nên được giữ làm bản gốc để chỉnh sửa."
    )

    document.add_heading("Cách chọn nhanh", level=1)
    quick_choices = [
        "Ảnh chụp để gửi hoặc đăng mạng xã hội: JPG hoặc JPEG.",
        "Ảnh trên website: WebP hoặc AVIF; có thể giữ JPG hoặc PNG làm phương án dự phòng.",
        "Logo và icon: SVG nếu cần phóng lớn; PNG nếu cần ảnh điểm ảnh có nền trong suốt.",
        "Ảnh động ngắn: WebP hoặc GIF.",
        "Chỉnh sửa chuyên nghiệp: RAW cho ảnh máy ảnh; PSD, AI, XCF hoặc KRA cho file thiết kế.",
        "Tài liệu gửi người khác hoặc in: PDF; quy trình in ảnh chuyên nghiệp có thể dùng TIFF.",
    ]
    for item in quick_choices:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.08
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5)

    document.add_heading("Bảng so sánh chi tiết", level=1)

    headers = ["Định dạng", "Loại và hỗ trợ", "Ưu điểm", "Hạn chế", "Nên dùng khi"]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    widths = [0.72, 1.35, 1.72, 1.58, 1.83]

    header_row = table.rows[0]
    repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, (cell, header, width) in enumerate(zip(header_row.cells, headers, widths)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell, top=110, start=105, bottom=110, end=105)
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(header)
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")

    for row_index, values in enumerate(FORMATS, start=1):
        row = table.add_row()
        prevent_row_split(row)
        fill = "EAF2F8" if row_index % 2 == 0 else "FFFFFF"
        for column_index, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.06
            lines = text.split("\n")
            for line_index, line in enumerate(lines):
                if line_index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                set_run_font(run, size=9.5, bold=(column_index == 0))

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    lead = note.add_run("Lưu ý  ")
    set_run_font(lead, size=10, bold=True)
    body = note.add_run(
        "Đổi phần mở rộng của tên file không chuyển đổi định dạng. Hãy dùng lệnh Save As, Export hoặc một công cụ chuyển đổi ảnh. "
        "Khi chỉnh sửa, nên giữ file gốc và xuất một bản riêng để chia sẻ."
    )
    set_run_font(body, size=10)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    label = footer_paragraph.add_run("Định dạng ảnh  •  Trang ")
    set_run_font(label, size=9, color="666666")
    add_page_number(footer_paragraph)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = create_document()
    print(path)
